#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
template_multinivel.py — Identidade visual dos relatórios Multinível.

Reproduz o relatório-modelo aprovado pelo Julio em 23/10/2026
("Relatório de Análise de Dados Cristiane 17.10 (teste cabeçalho e rodapé).docx"):

  - CABEÇALHO: faixa clara de largura total no topo de cada página (cabecalho-multinivel.png:
    lockup oficial positivo da Multinível à esquerda e o eyebrow "ANÁLISE DE DADOS ·
    PSICOMETRIA · DASHBOARDS" à direita), 2 cm de altura, atrás do texto.
  - RODAPÉ: faixa verde-profunda de largura total no pé de cada página (rodape-multinivel.png:
    lockup negativo oficial, site e contatos, malha de nós), com o número da página em
    Arial 9 branco, alinhado à direita, por cima da faixa.
  - MARCA D'ÁGUA: símbolo rede-M em g-600 a 14%, nós vazados com sombra suave
    (marca-dagua-multinivel.png), 13 cm, centralizado atrás do texto, em todas as páginas.
  - Distâncias: cabeçalho a 1,0 cm e rodapé a 0,8 cm da borda (o número cai no meio da faixa).

O build_report.py chama aplicar_template() antes de salvar, então todo relatório gerado sai
assim. Para um .docx que já existe:

  python3 template_multinivel.py relatorio.docx            -> "relatorio (Multinível).docx"
  python3 template_multinivel.py relatorio.docx --no-lugar  (sobrescreve)
  opções: --saida CAMINHO, --sem-faixas, --sem-marca, --largura-marca-cm N, --pular-primeira

É idempotente (reaplicar não duplica nada). Os três PNGs ficam ao lado deste script.
Para PDF: exporte o .docx pelo Word (tudo vai junto); marca_dagua.py carimba só a marca d'água.
Fontes das faixas e gerador: "2. Multinível/1. Consultorias/2026/00 - Marca d'água Multinível/
Cabeçalho e rodapé (teste)/" (gerar_faixas.py, variante A).
"""
import argparse, os, subprocess, sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image

AQUI = os.path.dirname(os.path.abspath(__file__))
PNG_CABECALHO = os.path.join(AQUI, "cabecalho-multinivel.png")
PNG_RODAPE = os.path.join(AQUI, "rodape-multinivel.png")
PNG_MARCA = os.path.join(AQUI, "marca-dagua-multinivel.png")
NOME_CAB, NOME_ROD = "Cabeçalho Multinível", "Rodapé Multinível"
IDIOMA_SCRIPT = os.path.expanduser("~/.claude/scripts/idioma_docx.py")
NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
      'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')


def _anchor(rid, cx, cy, x, y, doc_id, nome, arq, rh):
    return f'''<w:r {NS}><w:rPr><w:noProof/></w:rPr><w:drawing>
<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="{rh}" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">
<wp:simplePos x="0" y="0"/>
<wp:positionH relativeFrom="page"><wp:posOffset>{x}</wp:posOffset></wp:positionH>
<wp:positionV relativeFrom="page"><wp:posOffset>{y}</wp:posOffset></wp:positionV>
<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>
<wp:docPr id="{doc_id}" name="{nome}"/><wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>
<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic>
<pic:nvPicPr><pic:cNvPr id="0" name="{arq}"/><pic:cNvPicPr/></pic:nvPicPr>
<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>
<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>
</pic:pic></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>'''


def _max_docpr_id(doc):
    m = 0
    for part in doc.part.package.iter_parts():
        el = getattr(part, "element", None)
        if el is None:
            continue
        for d in el.iter(qn("wp:docPr")):
            try:
                m = max(m, int(d.get("id")))
            except (TypeError, ValueError):
                pass
    return m


def _remover(parte, nome):
    for d in list(parte._element.iter(qn("wp:docPr"))):
        if d.get("name") == nome:
            r = d
            while r is not None and r.tag != qn("w:r"):
                r = r.getparent()
            if r is not None and r.getparent() is not None:
                r.getparent().remove(r)


def _faixa(parte, png, sec, no_pe, doc_id, nome, rh):
    """Imagem de largura total (proporção preservada), ancorada à página, atrás do texto."""
    with Image.open(png) as im:
        prop = im.height / im.width
    cx = int(sec.page_width)
    cy = int(round(cx * prop / 3600.0)) * 3600  # altura em centésimos de cm (2,00 cm no modelo)
    y = int(sec.page_height) - cy if no_pe else 0
    rid, _ = parte.part.get_or_add_image(png)
    _remover(parte, nome)
    p = parte.paragraphs[0] if parte.paragraphs else parte.add_paragraph()
    p._p.append(parse_xml(_anchor(rid, cx, cy, 0, y, doc_id, nome, os.path.basename(png), rh)))
    return p


RPR_NUM = ('<w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/><w:color w:val="FFFFFF"/>'
           '<w:sz w:val="18"/><w:szCs w:val="18"/></w:rPr>')


def _numero_pagina(p):
    """Número da página em Arial 9 branco, à direita. Usa campo complexo (fldChar): com
    fldSimple o Word recalcula o campo com a fonte do parágrafo (Times 12 preto), e o número
    some sobre a faixa escura."""
    for el in list(p._p):
        if el.tag == qn("w:fldSimple"):
            p._p.remove(el)
        elif el.tag == qn("w:r") and (el.find(qn("w:fldChar")) is not None or el.find(qn("w:instrText")) is not None
                                      or el.get(qn("w:rsidR")) == "4D4E0001"):
            p._p.remove(el)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pf = p.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.0; pf.first_line_indent = Cm(0)
    # marca de parágrafo com a mesma formatação (o Word usa como referência em alguns casos)
    pPr = p._p.get_or_add_pPr()
    velho = pPr.find(qn("w:rPr"))
    if velho is not None:
        pPr.remove(velho)
    pPr.append(parse_xml(RPR_NUM.replace("<w:rPr>", f"<w:rPr {NS}>", 1)))
    partes = ['<w:fldChar w:fldCharType="begin"/>', '<w:instrText xml:space="preserve"> PAGE </w:instrText>',
              '<w:fldChar w:fldCharType="separate"/>', '<w:t>1</w:t>', '<w:fldChar w:fldCharType="end"/>']
    for x in partes:
        p._p.append(parse_xml(f'<w:r {NS} w:rsidR="4D4E0001">{RPR_NUM}{x}</w:r>'))


def aplicar_template(doc, faixas=True, marca=True, largura_marca_cm=13.0, pular_primeira=False):
    """Aplica cabeçalho, rodapé (com nº de página) e marca d'água num Document aberto.
    Devolve um dicionário com o que foi aplicado; avisa (sem falhar) se faltar arquivo."""
    feito = {"faixas": 0, "marca": 0}
    if marca:
        if os.path.isfile(PNG_MARCA):
            if AQUI not in sys.path:
                sys.path.insert(0, AQUI)
            import marca_dagua
            feito["marca"] = marca_dagua.marcar_document(doc, PNG_MARCA, largura_marca_cm, pular_primeira)
        else:
            print("[AVISO] marca-dagua-multinivel.png ausente: relatório SEM marca d'água.")
    if faixas:
        falta = [n for n in (PNG_CABECALHO, PNG_RODAPE) if not os.path.isfile(n)]
        if falta:
            print("[AVISO] faixas ausentes (" + ", ".join(os.path.basename(f) for f in falta) + "): relatório SEM cabeçalho e rodapé.")
        else:
            nid = _max_docpr_id(doc) + 1
            par_impar = doc.settings.element.find(qn("w:evenAndOddHeaders")) is not None
            for i, sec in enumerate(doc.sections):
                if i > 0 and sec.header.is_linked_to_previous and sec.footer.is_linked_to_previous:
                    continue
                sec.footer_distance = Cm(0.8)
                if sec.header_distance > Cm(1.0):
                    sec.header_distance = Cm(1.0)
                primeira = sec.different_first_page_header_footer
                hs = [sec.header] + ([sec.first_page_header] if primeira else []) + ([sec.even_page_header] if par_impar else [])
                fs = [sec.footer] + ([sec.first_page_footer] if primeira else []) + ([sec.even_page_footer] if par_impar else [])
                for h in hs:
                    _faixa(h, PNG_CABECALHO, sec, False, nid, NOME_CAB, 251657217); nid += 1
                for f in fs:
                    _numero_pagina(_faixa(f, PNG_RODAPE, sec, True, nid, NOME_ROD, 251657218)); nid += 1
                feito["faixas"] += len(hs)
    if feito["marca"]:
        print(f"Marca d'água Multinível aplicada ({feito['marca']} cabeçalho(s), {largura_marca_cm:g} cm).")
    if feito["faixas"]:
        print(f"Cabeçalho e rodapé Multinível aplicados ({feito['faixas']} cabeçalho(s), com número de página).")
    return feito


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("arquivos", nargs="+", help=".docx")
    ap.add_argument("--saida"); ap.add_argument("--no-lugar", action="store_true")
    ap.add_argument("--sem-faixas", action="store_true"); ap.add_argument("--sem-marca", action="store_true")
    ap.add_argument("--largura-marca-cm", type=float, default=13.0); ap.add_argument("--pular-primeira", action="store_true")
    a = ap.parse_args()
    if a.saida and len(a.arquivos) > 1:
        sys.exit("--saida só vale para um arquivo de entrada.")
    for entrada in a.arquivos:
        base, ext = os.path.splitext(entrada)
        if ext.lower() != ".docx":
            print(f"ignorado (só .docx; para PDF, exporte pelo Word): {entrada}"); continue
        doc = Document(entrada)
        aplicar_template(doc, not a.sem_faixas, not a.sem_marca, a.largura_marca_cm, a.pular_primeira)
        saida = entrada if a.no_lugar else (a.saida or f"{base} (Multinível).docx")
        tmp = saida + ".tmp.docx"
        doc.save(tmp); os.replace(tmp, saida)
        if os.path.exists(IDIOMA_SCRIPT):
            subprocess.run([sys.executable, IDIOMA_SCRIPT, saida], check=False, capture_output=True)
        print(f"✓ {saida}")


if __name__ == "__main__":
    main()
