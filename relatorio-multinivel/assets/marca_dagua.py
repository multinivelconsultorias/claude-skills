#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Aplica a marca d'água da Multinível (símbolo rede-M em g-600 a 14%, nós vazados com sombra) em relatórios
.docx e .pdf, centralizada na página, atrás do texto, em todas as páginas.

Uso:
  python3 marca_dagua.py relatorio.docx                → "relatorio (marca d'água).docx"
  python3 marca_dagua.py relatorio.pdf                 → "relatorio (marca d'água).pdf"
  python3 marca_dagua.py a.docx b.pdf c.docx           → vários de uma vez

Opções:
  --saida CAMINHO     nome do arquivo de saída (só com um arquivo de entrada)
  --no-lugar          sobrescreve o próprio arquivo de entrada
  --largura-cm N      largura da marca na página (padrão 13 cm)
  --pular-primeira    deixa a primeira página (capa) sem a marca
  --imagem PNG        usa outro PNG (padrão: marca-dagua-multinivel.png ao lado do script)

No .docx a marca entra como imagem flutuante "atrás do texto" ancorada no
cabeçalho de cada seção (é como o próprio Word faz a marca d'água de imagem),
por isso aparece em todas as páginas e sobrevive à exportação para PDF.
No .pdf a imagem é desenhada por baixo do conteúdo de cada página.
O build_report.py importa marcar_document() daqui e já aplica a marca em todo
relatório gerado; este script serve para relatórios já existentes e para PDFs.
Requer python-docx, Pillow e PyMuPDF.
"""
import argparse, os, shutil, subprocess, sys, copy

AQUI = os.path.dirname(os.path.abspath(__file__))
IMAGEM_PADRAO = os.path.join(AQUI, 'marca-dagua-multinivel.png')
NOME_MARCA = "Marca d'água Multinível"
EMU_POR_CM = 360000
IDIOMA_SCRIPT = os.path.expanduser('~/.claude/scripts/idioma_docx.py')


# ----------------------------------------------------------------- .docx
def _anchor_xml(rid, cx, cy, doc_id, nome_arquivo):
    return f'''<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
  xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
  xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
  xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"
  xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <w:rPr><w:noProof/></w:rPr>
  <w:drawing>
    <wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251657216"
               behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">
      <wp:simplePos x="0" y="0"/>
      <wp:positionH relativeFrom="page"><wp:align>center</wp:align></wp:positionH>
      <wp:positionV relativeFrom="page"><wp:align>center</wp:align></wp:positionV>
      <wp:extent cx="{cx}" cy="{cy}"/>
      <wp:effectExtent l="0" t="0" r="0" b="0"/>
      <wp:wrapNone/>
      <wp:docPr id="{doc_id}" name="{NOME_MARCA}" descr="Símbolo da Multinível em marca d'água"/>
      <wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>
      <a:graphic>
        <a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">
          <pic:pic>
            <pic:nvPicPr><pic:cNvPr id="0" name="{nome_arquivo}"/><pic:cNvPicPr/></pic:nvPicPr>
            <pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>
            <pic:spPr>
              <a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>
              <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
            </pic:spPr>
          </pic:pic>
        </a:graphicData>
      </a:graphic>
    </wp:anchor>
  </w:drawing>
</w:r>'''


def _max_docpr_id(doc):
    from docx.oxml.ns import qn
    maior = 0
    for part in doc.part.package.iter_parts():
        el = getattr(part, 'element', None)
        if el is None:
            continue
        for d in el.iter(qn('wp:docPr')):
            try:
                maior = max(maior, int(d.get('id')))
            except (TypeError, ValueError):
                pass
    return maior


def _remover_marca_existente(header):
    from docx.oxml.ns import qn
    for d in list(header._element.iter(qn('wp:docPr'))):
        if d.get('name') == NOME_MARCA:
            run = d
            while run is not None and run.tag != qn('w:r'):
                run = run.getparent()
            if run is not None and run.getparent() is not None:
                run.getparent().remove(run)


def _copiar_conteudo(origem, destino):
    """Copia os parágrafos/tabelas de um cabeçalho ou rodapé para outro (vazio)."""
    for el in list(destino._element):
        if el.tag.endswith('}p') or el.tag.endswith('}tbl'):
            destino._element.remove(el)
    for el in origem._element:
        if el.tag.endswith('}p') or el.tag.endswith('}tbl'):
            destino._element.append(copy.deepcopy(el))


def _inserir_no_header(header, imagem, largura_cm, doc_id):
    from docx.oxml import parse_xml
    rid, img = header.part.get_or_add_image(imagem)
    cx = int(round(largura_cm * EMU_POR_CM))
    cy = int(round(cx * img.px_height / img.px_width))
    _remover_marca_existente(header)
    par = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    par._p.append(parse_xml(_anchor_xml(rid, cx, cy, doc_id, os.path.basename(imagem))))


def marcar_document(doc, imagem=IMAGEM_PADRAO, largura_cm=13.0, pular_primeira=False):
    """Insere a marca nos cabeçalhos de um Document (python-docx) já aberto.
    Usado pelo build_report.py antes de salvar. Devolve o nº de cabeçalhos marcados."""
    from docx.oxml.ns import qn
    par_impar = doc.settings.element.find(qn('w:evenAndOddHeaders')) is not None
    prox_id = _max_docpr_id(doc) + 1
    n = 0
    for i, sec in enumerate(doc.sections):
        if i > 0 and sec.header.is_linked_to_previous:
            continue  # herda o cabeçalho (e a marca) da seção anterior
        if pular_primeira and i == 0 and not sec.different_first_page_header_footer:
            # separa a 1ª página preservando o que já havia no cabeçalho/rodapé
            sec.different_first_page_header_footer = True
            _copiar_conteudo(sec.header, sec.first_page_header)
            _copiar_conteudo(sec.footer, sec.first_page_footer)
        alvos = [sec.header]
        if sec.different_first_page_header_footer and not (pular_primeira and i == 0):
            alvos.append(sec.first_page_header)
        if par_impar:
            alvos.append(sec.even_page_header)
        for h in alvos:
            _inserir_no_header(h, imagem, largura_cm, prox_id); prox_id += 1; n += 1
    return n


def marcar_docx(entrada, saida, imagem, largura_cm, pular_primeira):
    from docx import Document
    doc = Document(entrada)
    n = marcar_document(doc, imagem, largura_cm, pular_primeira)
    doc.save(saida)
    if os.path.exists(IDIOMA_SCRIPT):
        subprocess.run([sys.executable, IDIOMA_SCRIPT, saida], check=False, capture_output=True)
    return n


# ------------------------------------------------------------------ .pdf
def marcar_pdf(entrada, saida, imagem, largura_cm, pular_primeira):
    import fitz
    from PIL import Image
    with Image.open(imagem) as im:
        prop = im.height / im.width
    w = largura_cm * 72 / 2.54
    h = w * prop
    tmp = saida + '.tmp.pdf'
    shutil.copyfile(entrada, tmp)

    def aplicar(overlay):
        doc = fitz.open(tmp)
        xref, n = 0, 0
        for pno, page in enumerate(doc):
            if pular_primeira and pno == 0:
                continue
            pr = page.rect
            rect = fitz.Rect(pr.x0 + (pr.width - w) / 2, pr.y0 + (pr.height - h) / 2,
                             pr.x0 + (pr.width + w) / 2, pr.y0 + (pr.height + h) / 2)
            rect = rect * page.derotation_matrix  # páginas giradas
            xref = page.insert_image(rect, filename=imagem if not xref else None, xref=xref,
                                     overlay=overlay, keep_proportion=True)
            n += 1
        doc.save(saida, garbage=3, deflate=True)
        doc.close()
        return n

    # 1ª tentativa: por baixo do conteúdo. Se a página tiver um fundo pintado que
    # esconde a marca, refaz por cima (a marca é translúcida, então o texto segue legível).
    n = aplicar(overlay=False)
    antes = fitz.open(entrada); depois = fitz.open(saida)
    pg = 1 if (pular_primeira and len(antes) > 1) else 0
    a = antes[pg].get_pixmap(dpi=36).samples; b = depois[pg].get_pixmap(dpi=36).samples
    antes.close(); depois.close()
    if a == b:
        print('  aviso: o fundo da página escondia a marca; aplicada por cima do conteúdo.')
        n = aplicar(overlay=True)
    os.remove(tmp)
    return n


# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('arquivos', nargs='+', help='.docx e/ou .pdf')
    ap.add_argument('--saida', help='arquivo de saída (apenas com uma entrada)')
    ap.add_argument('--no-lugar', action='store_true', help='sobrescreve a entrada')
    ap.add_argument('--largura-cm', type=float, default=13.0, help='largura da marca (padrão 13 cm)')
    ap.add_argument('--pular-primeira', action='store_true', help='não marca a primeira página')
    ap.add_argument('--imagem', default=IMAGEM_PADRAO, help='PNG da marca (padrão: ao lado do script)')
    a = ap.parse_args()
    if a.saida and len(a.arquivos) > 1:
        sys.exit('--saida só vale para um arquivo de entrada.')
    if not os.path.exists(a.imagem):
        sys.exit(f'Imagem da marca não encontrada: {a.imagem}')
    for entrada in a.arquivos:
        if not os.path.exists(entrada):
            print(f'não encontrado: {entrada}'); continue
        base, ext = os.path.splitext(entrada)
        ext_l = ext.lower()
        if ext_l not in ('.docx', '.pdf'):
            print(f'ignorado (só .docx e .pdf): {entrada}'); continue
        saida = entrada if a.no_lugar else (a.saida or f"{base} (marca d'água){ext}")
        if a.no_lugar:
            # grava num temporário e substitui, para não corromper o original se algo falhar
            tmp_out = f'{base}.__marca__{ext}'
            fn = marcar_docx if ext_l == '.docx' else marcar_pdf
            n = fn(entrada, tmp_out, a.imagem, a.largura_cm, a.pular_primeira)
            os.replace(tmp_out, entrada)
        else:
            fn = marcar_docx if ext_l == '.docx' else marcar_pdf
            n = fn(entrada, saida, a.imagem, a.largura_cm, a.pular_primeira)
        unidade = 'cabeçalho(s)' if ext_l == '.docx' else 'página(s)'
        print(f'✓ {os.path.basename(saida)}  ({n} {unidade} marcado(s), {a.largura_cm:g} cm)')


if __name__ == '__main__':
    main()
