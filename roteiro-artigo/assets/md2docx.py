#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Conversor mínimo de Markdown para .docx (cabeçalhos, tabelas, listas, citações, negrito e itálico)."""
import sys, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Times New Roman"

def style_run(run, size=11, bold=False, italic=False, color=None):
    run.font.name = FONT; run.font.size = Pt(size)
    run.font.bold = bold; run.font.italic = italic
    if color: run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii','w:hAnsi','w:cs'): rf.set(qn(a), FONT)

TOKEN = re.compile(r'(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)')

def add_rich(par, text, size=11, base_bold=False, color=None):
    """Escreve texto com **negrito**, *itálico* e `código`."""
    for part in TOKEN.split(text):
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            style_run(par.add_run(part[2:-2]), size, bold=True, color=color)
        elif part.startswith('*') and part.endswith('*'):
            style_run(par.add_run(part[1:-1]), size, italic=True, color=color)
        elif part.startswith('`') and part.endswith('`'):
            style_run(par.add_run(part[1:-1]), size, italic=True, color=color)
        else:
            style_run(par.add_run(part), size, bold=base_bold, color=color)

def fmt(par, after=6, before=0, spacing=1.15, indent=None):
    pf = par.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE; pf.line_spacing = spacing
    pf.space_after = Pt(after); pf.space_before = Pt(before)
    if indent is not None: pf.left_indent = Cm(indent)

def cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None: b = OxmlElement('w:tcBorders'); tcPr.append(b)
    for edge in ('top','bottom','left','right'):
        spec = kw.get(edge); el = b.find(qn('w:'+edge))
        if el is None: el = OxmlElement('w:'+edge); b.append(el)
        if spec is None: el.set(qn('w:val'),'nil')
        else:
            el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4')
            el.set(qn('w:space'),'0'); el.set(qn('w:color'),'000000')

def split_row(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]

def main(src, out):
    lines = open(src, encoding='utf8').read().split('\n')
    doc = Document()
    doc.styles['Normal'].font.name = FONT; doc.styles['Normal'].font.size = Pt(11)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2.5)
    i = 0
    while i < len(lines):
        ln = lines[i]; s = ln.strip()
        if not s or set(s) == {'-'} and len(s) > 2:
            i += 1; continue
        # tabela
        if s.startswith('|') and i+1 < len(lines) and set(lines[i+1].strip()) <= set('|-: '):
            header = split_row(s); rows = []
            j = i+2
            while j < len(lines) and lines[j].strip().startswith('|'):
                rows.append(split_row(lines[j])); j += 1
            ncol = len(header)
            tbl = doc.add_table(rows=0, cols=ncol); tbl.autofit = True
            hr = tbl.add_row().cells
            for k in range(ncol):
                p = hr[k].paragraphs[0]; fmt(p, after=2, spacing=1.0)
                add_rich(p, header[k] if k < len(header) else '', 10, base_bold=True)
                cell_border(hr[k], top={}, bottom={})
            for ri, r in enumerate(rows):
                cs = tbl.add_row().cells; last = ri == len(rows)-1
                for k in range(ncol):
                    p = cs[k].paragraphs[0]; fmt(p, after=2, spacing=1.0)
                    add_rich(p, r[k] if k < len(r) else '', 10)
                    cell_border(cs[k], bottom={} if last else None)
            fmt(doc.add_paragraph(), after=8)
            i = j; continue
        # cabeçalhos
        m = re.match(r'^(#{1,4})\s+(.*)', s)
        if m:
            lvl = len(m.group(1)); txt = m.group(2)
            sizes = {1: 15, 2: 13, 3: 12, 4: 11}
            p = doc.add_paragraph(); fmt(p, after=6, before=14 if lvl > 1 else 0)
            if lvl == 1: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_rich(p, txt, sizes[lvl], base_bold=True,
                     color=RGBColor(0x2C,0x5F,0x6F) if lvl <= 2 else None)
            i += 1; continue
        # citação em bloco
        if s.startswith('>'):
            p = doc.add_paragraph(); fmt(p, after=8, indent=1.0)
            add_rich(p, s.lstrip('> ').strip(), 11)
            i += 1; continue
        # lista
        m = re.match(r'^[-*]\s+(.*)', s)
        if m:
            p = doc.add_paragraph(style='List Bullet'); fmt(p, after=3)
            add_rich(p, m.group(1), 11)
            i += 1; continue
        # parágrafo comum
        p = doc.add_paragraph(); fmt(p, after=8)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_rich(p, s, 11)
        i += 1
    doc.save(out); print("OK ->", out)

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
